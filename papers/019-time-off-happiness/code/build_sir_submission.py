from __future__ import annotations

import hashlib
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PAPER = Path(__file__).resolve().parents[1]
SOURCE = PAPER / "manuscript" / "SIR_SUBMISSION_BLINDED.md"
FIG_DIR = PAPER / "figures"
OUT = PAPER / "submission"

DOCX_NAME = "SIR_SUBMISSION_BLINDED.docx"
PDF_NAME = "SIR_REVIEWER_MANUSCRIPT_BLINDED.pdf"
QA_NAME = "SIR_SUBMISSION_QA.md"
MANIFEST_NAME = "SIR_SUBMISSION_PACKAGE_MANIFEST.md"
REVIEW_NAME = "ANONYMOUS_REVIEW_PACKAGE_MANIFEST.md"

FIGS = [
    ("fig1_eight_event_refresh.svg", "Eight-event WHR2024 refreshed donor-adjusted full-post Life Ladder gaps; Bahrain is a large positive outlier while six of eight event means are negative."),
    ("fig2_legal_isolation_funnel.svg", "Legal treatment-isolation funnel showing legacy World Bank leave jumps narrowing to zero new clean holdouts and modern legal snapshots yielding Israel 2016 as a clean holdout and Mexico 2023 as a future holdout."),
    ("fig3_israel_event_time.svg", "Israel donor-adjusted Life Ladder event-time gaps under the strict 115-country donor pool; 2016 is excluded as a transition year and the full-post mean is approximately zero."),
    ("fig4_israel_reference_sensitivity.svg", "Israel reference-baseline sensitivity showing the frozen 2015 baseline near zero while earlier and multi-year post-outcome diagnostic baselines are negative."),
]
IDENTITY_TERMS = ["Cochrane", "Kang", "Cunyi", "ARIS4C", "github.com/CochraneK", "CochraneK"]


def run(args: list[str], *, cwd: Path | None = None, capture: bool = False) -> str:
    p = subprocess.run(
        args,
        cwd=str(cwd) if cwd else None,
        check=True,
        text=True,
        stdout=subprocess.PIPE if capture else None,
        stderr=subprocess.STDOUT if capture else None,
    )
    return p.stdout if capture else ""


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def set_style_font(style, name: str, size: float | None = None, bold: bool | None = None):
    style.font.name = name
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run0 = paragraph.add_run()
    for tag, attrs, text0 in [
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, " PAGE "),
        ("w:fldChar", {"w:fldCharType": "separate"}, None),
        ("w:t", {}, "1"),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ]:
        node = OxmlElement(tag)
        for k, v in attrs.items():
            node.set(qn(k), v)
        if text0 is not None:
            node.text = text0
        run0._r.append(node)


def build_reference_docx(path: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    set_style_font(doc.styles["Normal"], "Times New Roman", 12)
    doc.styles["Normal"].paragraph_format.line_spacing = 2
    doc.styles["Normal"].paragraph_format.space_after = Pt(0)
    for name, size in [("Heading 1", 14), ("Heading 2", 13), ("Heading 3", 12)]:
        set_style_font(doc.styles[name], "Times New Roman", size, True)
        doc.styles[name].paragraph_format.space_before = Pt(8)
        doc.styles[name].paragraph_format.space_after = Pt(4)
    for name in ["Body Text", "First Paragraph"]:
        if name in doc.styles:
            set_style_font(doc.styles[name], "Times New Roman", 12)
            doc.styles[name].paragraph_format.line_spacing = 2
            doc.styles[name].paragraph_format.space_after = Pt(0)
    add_page_field(sec.footer.paragraphs[0])
    doc.save(path)


def prepare_markdown(src: str, png_dir: Path) -> str:
    src = src.replace("\n---\n", "\n", 1)
    pat = re.compile(
        r"\n\$\s*\n\\{1,2}operatorname\{Gap\}_\{it\}\s*=\s*\(Y_\{it\}\s*-\s*Y_\{i,ref\}\)\s*-\s*"
        r"\\{1,2}frac\{1\}\{N_\{it\}\}\\{1,2}sum_\{d\s*\\{1,2}in\s*D_i\}\(Y_\{dt\}\s*-\s*Y_\{d,ref\}\),?\s*\n\$\s*\n"
    )
    src, n = pat.subn("\n\nEQGAPPLACEHOLDER\n\n", src, count=1)
    if n != 1:
        pat2 = re.compile(r"\n\$\$.*?\\{1,2}operatorname\{Gap\}_\{it\}.*?\$\$\s*\n", re.S)
        src, n = pat2.subn("\n\nEQGAPPLACEHOLDER\n\n", src, count=1)
    if n != 1:
        raise RuntimeError("Estimator equation block not found exactly once")
    for fname, _ in FIGS:
        old = f"../figures/{fname}"
        new = str((png_dir / fname.replace(".svg", ".png")).resolve())
        src = src.replace(f"![]({old})", f"![]({new}){{width=6.3in}}")
    return src


def clear_runs(paragraph):
    for r in list(paragraph.runs):
        paragraph._p.remove(r._r)


def add_gap_equation(paragraph):
    clear_runs(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)

    def add(text: str, *, sub=False, italic=False):
        r = paragraph.add_run(text)
        r.font.name = "Cambria Math"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria Math")
        r.font.size = Pt(12)
        r.font.subscript = sub
        r.font.italic = italic

    add("Gap"); add("it", sub=True, italic=True); add(" = (")
    add("Y", italic=True); add("it", sub=True, italic=True); add(" − ")
    add("Y", italic=True); add("i,ref", sub=True, italic=True); add(") − (1/")
    add("N", italic=True); add("it", sub=True, italic=True); add(") Σ")
    add("d∈Dᵢ", sub=True, italic=True); add(" (")
    add("Y", italic=True); add("dt", sub=True, italic=True); add(" − ")
    add("Y", italic=True); add("d,ref", sub=True, italic=True); add(")")


def postprocess_docx(path: Path):
    doc = Document(path)
    cp = doc.core_properties
    cp.author = ""; cp.last_modified_by = ""; cp.comments = ""; cp.category = ""
    cp.content_status = ""; cp.identifier = ""; cp.language = "en"
    cp.subject = "Blinded manuscript for peer review"
    cp.title = "Statutory Paid Annual Leave and National Life Evaluation"
    cp.keywords = "paid annual leave; life evaluation; social indicators; legal epidemiology; causal inference; quality of life"

    for style in doc.styles:
        try:
            style.font.color.rgb = RGBColor(0, 0, 0)
            if style.font.name:
                style._element.rPr.rFonts.set(qn("w:eastAsia"), style.font.name)
        except Exception:
            pass
    for p in doc.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor(0, 0, 0)
            if not r.font.name:
                r.font.name = "Times New Roman"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    for idx, size in [(0, 16), (1, 13)]:
        p = doc.paragraphs[idx]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1
        p.paragraph_format.space_after = Pt(6)
        for r in p.runs:
            r.bold = True
            r.font.size = Pt(size)
            r.font.name = "Times New Roman"

    equation_hits = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt == "EQGAPPLACEHOLDER":
            add_gap_equation(p)
            equation_hits += 1
        if txt == "1. Introduction":
            p.paragraph_format.page_break_before = True
        if txt.startswith("Fig. "):
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(6)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)
        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(2)
    if equation_hits != 1:
        raise RuntimeError(f"Expected one equation placeholder, found {equation_hits}")

    nodes = doc._element.xpath(".//wp:docPr")
    if len(nodes) != len(FIGS):
        raise RuntimeError(f"Expected {len(FIGS)} images, found {len(nodes)}")
    for idx, (node, (_, alt)) in enumerate(zip(nodes, FIGS), start=1):
        node.set("descr", alt)
        node.set("title", alt.split(";")[0])
        node.set("name", f"Figure {idx}")

    # Pandoc writes the absolute temporary image path into pic:cNvPr/@descr.
    # That path can expose the local project/workspace name even when the
    # manuscript text and core properties are blinded. Scrub the picture-level
    # non-visual properties as well as wp:docPr before saving.
    pic_nodes = doc._element.xpath(".//pic:cNvPr")
    if len(pic_nodes) != len(FIGS):
        raise RuntimeError(f"Expected {len(FIGS)} picture metadata nodes, found {len(pic_nodes)}")
    for idx, (node, (_, alt)) in enumerate(zip(pic_nodes, FIGS), start=1):
        node.set("descr", alt)
        node.set("name", f"Figure {idx}")

    doc.save(path)


def privacy_scrub(path: Path):
    tmp = path.with_suffix(".scrub.docx")
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                root = ET.fromstring(data)
                ns = {
                    "dc": "http://purl.org/dc/elements/1.1/",
                    "cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties",
                }
                for xp in ["dc:creator", "cp:lastModifiedBy"]:
                    el = root.find(xp, ns)
                    if el is not None:
                        el.text = ""
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            zout.writestr(item, data)
    tmp.replace(path)


def pdf_pages(pdf: Path) -> int:
    out = run(["pdfinfo", str(pdf)], capture=True)
    m = re.search(r"^Pages:\s+(\d+)", out, re.M)
    if not m:
        raise RuntimeError("Could not determine PDF page count")
    return int(m.group(1))


def inspect_docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        parts = []
        for name in ["word/document.xml", "docProps/core.xml", "docProps/app.xml"]:
            if name in z.namelist():
                parts.append(z.read(name).decode("utf-8", errors="ignore"))
        return "\n".join(parts)


def identity_hits(text: str) -> list[str]:
    low = text.lower()
    return [term for term in IDENTITY_TERMS if term.lower() in low]


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for exe in ["pandoc", "libreoffice", "inkscape", "pdftotext", "pdfinfo"]:
        if shutil.which(exe) is None:
            raise RuntimeError(f"Required executable missing: {exe}")
    if not SOURCE.exists():
        raise RuntimeError(f"Missing manuscript source: {SOURCE}")

    with tempfile.TemporaryDirectory(prefix="sir_build_", dir=str(PAPER)) as td0:
        td = Path(td0)
        png_dir = td / "figures"
        png_dir.mkdir()
        for fname, _ in FIGS:
            src = FIG_DIR / fname
            dst = png_dir / fname.replace(".svg", ".png")
            run(["inkscape", str(src), "--export-type=png", f"--export-filename={dst}", "--export-width=2400"])

        build_md = td / "SIR_SUBMISSION_BLINDED_BUILD.md"
        build_md.write_text(prepare_markdown(SOURCE.read_text(encoding="utf-8"), png_dir), encoding="utf-8")
        ref = td / "reference.docx"
        build_reference_docx(ref)
        draft = td / DOCX_NAME
        run(["pandoc", str(build_md), "--reference-doc", str(ref), "-o", str(draft)])
        postprocess_docx(draft)
        privacy_scrub(draft)

        final_docx = OUT / DOCX_NAME
        shutil.copy2(draft, final_docx)
        run(["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(td), str(final_docx)])
        exported = td / final_docx.with_suffix(".pdf").name
        final_pdf = OUT / PDF_NAME
        shutil.copy2(exported, final_pdf)

    pdf_text_path = OUT / ".sir_pdf_text.txt"
    run(["pdftotext", str(OUT / PDF_NAME), str(pdf_text_path)])
    pdf_text = pdf_text_path.read_text(encoding="utf-8", errors="ignore")
    pdf_text_path.unlink(missing_ok=True)
    docx_xml = inspect_docx_text(OUT / DOCX_NAME)
    hits = sorted(set(identity_hits(docx_xml) + identity_hits(pdf_text)))
    with zipfile.ZipFile(OUT / DOCX_NAME) as z:
        image_count = sum(1 for n in z.namelist() if n.startswith("word/media/") and not n.endswith("/"))
    page_count = pdf_pages(OUT / PDF_NAME)
    fig_caption_count = len(re.findall(r"Fig\.\s*[1-4]", pdf_text))
    equation_ok = "Gap" in pdf_text and "operatorname" not in pdf_text and "EQGAPPLACEHOLDER" not in pdf_text
    qa_pass = not hits and image_count == 4 and fig_caption_count >= 4 and equation_ok and page_count > 0

    qa = f"""# SIR blinded document-package QA

- Status: **{"PASS" if qa_pass else "FAIL"}**
- Source: manuscript/SIR_SUBMISSION_BLINDED.md
- Editable DOCX: {DOCX_NAME}
- Reviewer PDF: {PDF_NAME}
- PDF pages: **{page_count}**
- Embedded figures: **{image_count}/4**
- Figure-caption references detected in extracted PDF text: **{fig_caption_count}**
- Estimator equation rendered without LaTeX control-string leakage: **{equation_ok}**
- Identity scan terms found in DOCX XML/PDF text: **{", ".join(hits) if hits else "none"}**
- DOCX SHA-256: {sha256(OUT / DOCX_NAME)}
- PDF SHA-256: {sha256(OUT / PDF_NAME)}

## Visual QA contract

The Word and PDF are generated from the same blinded Markdown source. The controller separately records page-by-page visual QA of pagination, equation rendering, figure/caption pairs, reference completion, and clipping.

## Anonymity boundary

No author name, repository owner, project identifier, or identifying repository URL is permitted in the blinded manuscript files. Author metadata and institution-specific declarations remain outside this package.
"""
    (OUT / QA_NAME).write_text(qa, encoding="utf-8")

    review_manifest = """# Anonymous reviewer package manifest

**Status:** reviewer-facing delivery is defined as **Online Resource 1**, a de-identified replication ZIP uploaded directly through the journal submission system.

The identifying development repository must not be linked from the blinded manuscript or reviewer files.

## Online Resource 1

Expected file: `SIR_Online_Resource_1_Anonymous_Replication.zip`

The archive is built separately by `build_sir_anonymous_review_bundle.py` and must pass its identity scan and number-lock reproduction gate before portal submission.

## Include

- frozen event and donor registries;
- author-derived machine-readable result tables;
- one-command reproduction code;
- source-provenance and licensing notes.

## Do not redistribute

- raw Gallup / World Happiness Report files where source terms prohibit redistribution;
- Equal Futures raw data where source terms restrict redistribution;
- author identity, affiliation, email, ORCID, private notes, chat logs, or identifying repository URLs.
"""
    (OUT / REVIEW_NAME).write_text(review_manifest, encoding="utf-8")

    fig_lines = "\n".join(
        f"- figures/{fname} SHA-256 {sha256(FIG_DIR / fname)}" for fname, _ in FIGS
    )
    manifest = f"""# SIR submission package manifest

**Package state:** document generation and machine QA {"PASS" if qa_pass else "FAIL"}; portal submission additionally requires a PASS Online Resource 1 anonymous replication bundle and author-side metadata/declarations.

## Frozen manuscript inputs

- Blinded Markdown source SHA-256: {sha256(SOURCE)}
{fig_lines}

## Generated reviewer documents

- {DOCX_NAME} — SHA-256 {sha256(OUT / DOCX_NAME)}
- {PDF_NAME} — SHA-256 {sha256(OUT / PDF_NAME)}
- PDF page count: **{page_count}**
- Machine QA: {QA_NAME}
- Anonymous review manifest: {REVIEW_NAME}

## Scientific lock

Packaging does not reopen outcomes, event admission, estimator selection, or the locked interpretation. Positive/negative affect remain closed for this paper.

## Pending external/author-only fields

- PASS `SIR_Online_Resource_1_Anonymous_Replication.zip` reviewer bundle;
- affiliation and corresponding-author email;
- ORCID if used;
- funding declaration;
- competing-interest declaration;
- institution-specific ethics/exemption determination;
- final journal-portal metadata validation.

Until those fields are supplied, this is a submission document package, not a completed portal submission.
"""
    (OUT / MANIFEST_NAME).write_text(manifest, encoding="utf-8")

    if not qa_pass:
        raise RuntimeError("SIR document package QA failed")
    print(qa)


if __name__ == "__main__":
    main()
